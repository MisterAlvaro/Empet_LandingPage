import sys
from pathlib import Path
from typing import List, Tuple

from pptx import Presentation
from pptx.util import Pt


def read_text(path: Path) -> str:
	return path.read_text(encoding="utf-8")


def split_into_slides(markdown: str) -> List[str]:
	lines = markdown.splitlines()
	slides: List[str] = []
	current: List[str] = []
	for line in lines:
		if line.strip() == "---":
			if current:
				slides.append("\n".join(current).strip())
				current = []
			continue
		current.append(line)
	if current:
		slides.append("\n".join(current).strip())
	return slides


def parse_slide(slide_md: str) -> Tuple[str, List[str], List[str]]:
	"""Return (title, bullets, notes)."""
	lines = [l for l in slide_md.splitlines() if l.strip() != ""]
	title = ""
	bullets: List[str] = []
	notes: List[str] = []

	# Find title (first markdown heading line starting with ### or ## or #)
	for line in lines:
		ls = line.lstrip()
		if ls.startswith("### "):
			title = ls[4:].strip()
			break
		elif ls.startswith("## "):
			title = ls[3:].strip()
			break
		elif ls.startswith("# "):
			title = ls[2:].strip()
			break

	# Identify position of "Guion sugerido:" to split body vs notes
	try:
		idx_notes = next(i for i, l in enumerate(lines) if l.lower().startswith("guion sugerido"))
	except StopIteration:
		idx_notes = -1

	# Collect bullets before notes section
	content_lines = lines if idx_notes == -1 else lines[:idx_notes]
	for line in content_lines:
		ls = line.lstrip()
		if ls.startswith("- "):
			bullets.append(ls[2:].strip())

	# Collect notes after notes header
	if idx_notes != -1:
		for line in lines[idx_notes + 1 :]:
			text = line.strip()
			if text.startswith("- "):
				text = text[2:].strip()
			# Strip surrounding quotes if present
			if (text.startswith('"') and text.endswith('"')) or (text.startswith("'") and text.endswith("'")):
				text = text[1:-1]
			if text:
				notes.append(text)

	return title, bullets, notes


def add_slide(prs: Presentation, title: str, bullets: List[str], notes: List[str]) -> None:
	layout = prs.slide_layouts[1]  # Title and Content
	slide = prs.slides.add_slide(layout)
	slide.shapes.title.text = title or ""
	body = slide.shapes.placeholders[1].text_frame
	body.clear()
	if bullets:
		# First bullet sets text
		body.text = bullets[0]
		p = body.paragraphs[0]
		p.font.size = Pt(20)
		# Remaining bullets
		for b in bullets[1:]:
			para = body.add_paragraph()
			para.text = b
			para.level = 0
			para.font.size = Pt(20)
	else:
		body.text = ""

	# Speaker notes
	notes_slide = slide.notes_slide
	notes_tf = notes_slide.notes_text_frame
	notes_tf.clear()
	if notes:
		notes_tf.text = notes[0]
		for n in notes[1:]:
			p = notes_tf.add_paragraph()
			p.text = n
	else:
		notes_tf.text = ""


def convert(md_path: Path, out_path: Path) -> None:
	markdown = read_text(md_path)
	slides_md = split_into_slides(markdown)
	prs = Presentation()
	for slide_md in slides_md:
		title, bullets, notes = parse_slide(slide_md)
		add_slide(prs, title, bullets, notes)
	prs.save(str(out_path))


def main() -> None:
	if len(sys.argv) < 3:
		print("Usage: python md_to_pptx.py <input.md> <output.pptx>")
		sys.exit(1)
	md_path = Path(sys.argv[1]).resolve()
	out_path = Path(sys.argv[2]).resolve()
	convert(md_path, out_path)
	print(f"Wrote {out_path}")


if __name__ == "__main__":
	main()